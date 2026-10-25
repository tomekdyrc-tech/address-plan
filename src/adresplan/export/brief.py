"""Compact reports for recipients who need allocations rather than engine details."""

from io import BytesIO

from adresplan.application.messages import FINDING_MESSAGES
from adresplan.core.models import Project
from adresplan.export.common import create_report, document_details, ipv6_description
from adresplan.export.markdown import _table, markdown_text
from adresplan.i18n import tr, tr_all


def brief_data(
    project: Project,
) -> tuple[tuple[str, ...], list[tuple[object, ...]], list[str], str]:
    report = create_report(project)
    # Polish header ids; each output translates them when rendering.
    headers: tuple[str, ...] = ("Lokalizacja", "Budynek", "Sieć", "VLAN", "IPv4", "Brama")
    if project.uses_ipv6:
        headers += ("IPv6",)
    rows = []
    for r in report.rows:
        row: tuple[object, ...] = (
            r.site.name,
            r.building.name,
            r.segment.name,
            r.segment.vlan_id,
            r.v4,
            ", ".join(map(str, r.reservation.gateway_addresses)) if r.reservation else "",
        )
        if project.uses_ipv6:
            row += (r.v6,)
        rows.append(row)
    names = {n.id: n.name for n in project.nodes}
    warnings = [
        f"{names.get(f.node_id or '', project.name)}: "
        + tr(
            FINDING_MESSAGES.get(
                f.code,
                FINDING_MESSAGES.get(f.code.removeprefix("v6_"), "Sprawdź reguły planowania."),
            )
        )
        for f in report.plan.findings
        if f.level != "info"
    ]
    return headers, rows, warnings, report.status


def dumps_brief_markdown(project: Project) -> str:
    headers, rows, warnings, status = brief_data(project)
    sections = [
        "# " + tr("Plan adresacji — {name}", name=markdown_text(project.name)),
        "**" + tr("Status: {status}. Raport skrócony.", status=status) + "**",
        _table((tr("Dokumentacja"), tr("Wartość")), document_details(project)),
        f"IPv4: {project.v4_pool}. IPv6: {markdown_text(ipv6_description(project))}.",
        _table(tr_all(*headers), rows),
    ]
    sections.extend(f"- {markdown_text(w)}" for w in warnings)
    return "\n\n".join(sections) + "\n"


def brief_office(project: Project, format: str) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm
    from openpyxl import Workbook

    from adresplan.export.office import _sheet, office_text
    from adresplan.export.office import _table as docx_table

    headers, rows, warnings, status = brief_data(project)
    output = BytesIO()
    if format == "xlsx":
        book = Workbook()
        book.remove(book["Sheet"])
        _sheet(book, "Plan", project, status, headers, rows)
        _sheet(
            book,
            "Dokumentacja",
            project,
            status,
            ("Parametr", "Wartość"),
            (
                *document_details(project),
                ("IPv6", ipv6_description(project)),
                *((tr("Uwaga"), w) for w in warnings),
            ),
        )
        book.save(output)
    else:
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21)
        section.left_margin = section.right_margin = Cm(1.5)
        document.add_paragraph(office_text(project.name), "Title")
        document.add_paragraph(tr("Status: {status}. Raport skrócony.", status=status))
        for label, value in document_details(project):
            document.add_paragraph(office_text(f"{label}: {value}"))
        document.add_paragraph(
            office_text(f"IPv4: {project.v4_pool}. IPv6: {ipv6_description(project)}.")
        )
        docx_table(document, headers, rows, tuple(26 / len(headers) for _ in headers))
        for warning in warnings:
            document.add_paragraph(office_text(warning))
        document.save(output)
    return output.getvalue()

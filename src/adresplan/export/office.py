"""Portable Office reports generated from the validated shared report model."""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import fields
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from adresplan.application.console import format_finding
from adresplan.application.project_json import InputIssue, ProjectJSONError
from adresplan.core.models import IPv6PoolMetrics, PoolMetrics, Project
from adresplan.export.common import address_range, create_report, document_details, percentage
from adresplan.export.csv_ import CSV_COLUMNS, report_records

_IDENTITY_COLUMNS = ("Lokalizacja", "Budynek", "Segment")
XLSX_COLUMNS = (*_IDENTITY_COLUMNS, *(c for c in CSV_COLUMNS if c not in _IDENTITY_COLUMNS))


def office_text(value: object) -> str:
    """Keep text literal; expose XML-forbidden controls instead of dropping data."""
    text = "—" if value is None or value == "" else str(value)
    return re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", lambda match: f"\\u{ord(match[0]):04x}", text
    )


def _sheet(
    book: Workbook,
    name: str,
    project: Project,
    status: str,
    headers: Iterable[str],
    rows: Iterable[Iterable[object]],
) -> Worksheet:
    sheet: Worksheet = book.create_sheet(name)
    header_values = list(headers)
    values: list[list[object]] = [list(header_values), *[list(row) for row in rows]]
    count = len(values[0])
    sheet.append([office_text(project.name)])
    sheet.append([f"Status: {status} · Źródło: projekt {office_text(project.id)}"])
    sheet.append(
        ["Migawka planu. Zmiany adresacji wprowadź w Address Plan i wyeksportuj ponownie."]
    )
    for cells in values:
        if any(len(office_text(value)) > 32767 for value in cells):
            raise ProjectJSONError(
                (
                    InputIssue(
                        (),
                        "xlsx_cell_limit",
                        "Komórka XLSX przekracza limit 32767 znaków. "
                        "Skróć uwagi albo użyj DOCX/JSON.",
                    ),
                )
            )
        sheet.append([office_text(value) for value in cells])
    # Every cell is a literal string, even '=...', errors, big IPv6 counts and IPs.
    for row_number, cell_row in enumerate(sheet, 1):
        for cell in cell_row:
            if cell.value is not None:
                cell.data_type = "s"
                cell.number_format = "@"
            cell.font = Font(name="Calibri", size=11, color="172D3A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_number >= 5 and row_number % 2:
                cell.fill = PatternFill("solid", fgColor="F0F5F7")
    for row_number in (1, 2, 3):
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=count
        )
        sheet.row_dimensions[row_number].height = 24 if row_number != 3 else 32
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color="172D3A")
    for cell in sheet[4]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B49")
    sheet.row_dimensions[4].height = 34
    for index, header in enumerate(header_values, 1):
        width = 48 if any(w in header for w in ("Uwagi", "Walidacja", "Komunikat")) else 24
        if header in ("VLAN", "Rola", "Błędy", "Ostrzeżenia"):
            width = 14
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index, value_row in enumerate(values[1:], 5):
        lines = max(
            (len(office_text(v)) // 35 + office_text(v).count("\n") + 1 for v in value_row),
            default=1,
        )
        sheet.row_dimensions[index].height = min(300, max(32, lines * 15))
    sheet.freeze_panes = "D5" if name == "Plan" else "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(count)}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:4"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if count > 12 else sheet.PAPERSIZE_A4
    # Wide planning tables span pages horizontally; never shrink 32 columns to unreadable text.
    assert sheet.sheet_properties.pageSetUpPr is not None
    sheet.sheet_properties.pageSetUpPr.fitToPage = count <= 12
    sheet.page_setup.fitToWidth = 0 if count > 12 else 1
    sheet.page_setup.fitToHeight = 0
    if count > 12:
        sheet.page_setup.scale = 100
        sheet.print_title_cols = "A:C"
    sheet.print_area = sheet.dimensions
    return sheet


def dumps_xlsx(project: Project) -> bytes:
    report = create_report(project)
    book = Workbook()
    book.remove(book["Sheet"])
    records = list(report_records(project))
    _sheet(
        book,
        "Plan",
        project,
        report.status,
        XLSX_COLUMNS,
        ([record.get(column, "") for column in XLSX_COLUMNS] for record in records),
    )
    names = {n.id: n.name for n in project.nodes}
    summary: list[tuple[object, ...]] = [
        ("Pula IPv4", project.v4_pool, "", "", "", ""),
        ("Prefiks IPv6", project.ipv6_description, "", "", "", ""),
    ]
    summary.extend((label, value, "", "", "", "") for label, value in document_details(project))
    metrics: tuple[PoolMetrics | IPv6PoolMetrics, ...] = (
        *report.plan.v4_metrics,
        *report.plan.v6_metrics,
    )
    summary.extend(
        ("Ustawienie: " + f.name, getattr(project.settings, f.name), "", "", "", "")
        for f in fields(project.settings)
    )
    summary.extend(
        (
            names.get(m.parent_id or "", project.name),
            m.network,
            percentage(m.utilization),
            m.allocated_addresses,
            m.free_addresses,
            m.largest_free_block,
        )
        for m in metrics
    )
    _sheet(
        book,
        "Podsumowanie",
        project,
        report.status,
        (
            "Zakres / parametr",
            "Pula / wartość",
            "Wykorzystanie",
            "Zajęte adresy",
            "Wolne adresy",
            "Największy wolny blok",
        ),
        summary,
    )
    _sheet(
        book,
        "Wolne bloki",
        project,
        report.status,
        ("Rodzic", "Rodzaj puli", "CIDR"),
        (
            (names.get(b.parent_id or "", project.name), b.pool_kind or "IPv4", b.network)
            for b in report.plan.free_blocks
        ),
    )
    _sheet(
        book,
        "Walidacja",
        project,
        report.status,
        ("Poziom", "Kod", "Element", "Komunikat"),
        (
            (f.level, f.code, names.get(f.node_id or "", project.name), format_finding(f, project))
            for f in report.plan.findings
        ),
    )
    if project.active_blocked_subnets:
        _sheet(
            book,
            "Blokady",
            project,
            report.status,
            ("Pula", "CIDR", "Opis"),
            (
                (names.get(b.parent_id or "", project.name), b.network, b.note)
                for b in project.active_blocked_subnets
            ),
        )
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def _table(
    document: Any,
    headers: tuple[str, ...],
    rows: Iterable[Iterable[object]],
    widths: tuple[float, ...],
) -> None:
    if document.paragraphs:
        document.paragraphs[-1].paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + name)
        for key, border_value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), border_value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Cm(width)
    for i, values in enumerate((headers, *[tuple(row) for row in rows])):
        row = table.rows[0] if i == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, values, widths, strict=True):
            cell.width = Cm(width)
            cell.vertical_alignment = 1
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "173B49" if i == 0 else "F0F5F7" if i % 2 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            cell.text = office_text(value)
            for paragraph in cell.paragraphs:
                compact = headers[0] in ("Rodzic", "Poziom")
                paragraph.paragraph_format.space_after = Pt(2 if compact else 5)
                paragraph.paragraph_format.space_before = Pt(2 if compact else 5)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5 if compact else 10)
                    run.font.bold = i == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "000000")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def dumps_docx(project: Project) -> bytes:
    report = create_report(project)
    document = Document()
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.left_margin = section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(2)
    section.footer_distance = Cm(0.7)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.add_paragraph("Plan adresacji sieci", "Title")
    document.add_paragraph(office_text(project.name), "Subtitle")
    for label, value in document_details(project):
        document.add_paragraph(office_text(f"{label}: {value}"))
    document.add_paragraph(
        f"Status planu: {report.status}. Dokument przedstawia lokalizacje, budynki, "
        "segmenty i rezerwacje. Wynik niekompletny zawiera wyłącznie udane przydziały; "
        "przed wdrożeniem popraw błędy wymienione w walidacji."
    )
    document.add_heading("Założenia", 1)
    document.add_paragraph(
        f"Pula IPv4: {project.v4_pool}. Prefiks IPv6: {project.ipv6_description}. "
        f"Rodzaj IPv6: {project.v6_kind}. Hosty nie obejmują osobnych bram i rezerwy statycznej. "
        "Zapas hostów jest zaokrąglany w górę. P2P i loopback mają osobne reguły."
    )
    _table(
        document,
        ("Ustawienie", "Wartość", "Ustawienie", "Wartość"),
        (
            (
                f.name,
                getattr(project.settings, f.name),
                g.name if g else "",
                getattr(project.settings, g.name) if g else "",
            )
            for f, g in zip(
                fields(project.settings)[::2], (*fields(project.settings)[1::2], None), strict=False
            )
        ),
        (7.5, 5.5, 7.5, 5.8),
    )
    names = {n.id: n.name for n in project.nodes}
    allocations4 = {a.node_id: a.network for a in report.plan.v4_allocations}
    allocations6 = {a.node_id: a.network for a in report.plan.v6_allocations}
    for site in report.sites:
        document.add_heading("Lokalizacja " + office_text(site.name), 1)
        document.add_paragraph(
            f"IPv4: {allocations4.get(site.id, '—')} · IPv6: {allocations6.get(site.id, '—')}"
        )
        if site.notes:
            document.add_paragraph(office_text(site.notes))
        for building in report.buildings:
            if building.parent_id != site.id:
                continue
            document.add_heading("Budynek " + office_text(building.name), 2)
            document.add_paragraph(
                f"IPv4: {allocations4.get(building.id, '—')} · "
                f"IPv6: {allocations6.get(building.id, '—')}"
            )
            if building.notes:
                document.add_paragraph(office_text(building.notes))
            rows = [r for r in report.rows if r.building.id == building.id]
            if not rows:
                document.add_paragraph("Brak segmentów.")
                continue
            _table(
                document,
                (
                    "Segment i VLAN",
                    "Hosty / zapas",
                    "IPv4 i maska",
                    "Brama IPv4",
                    "Statyczne / DHCP",
                ),
                (
                    (
                        f"{r.segment.name}\nVLAN {r.segment.vlan_id or '—'} · {r.segment.role}",
                        f"{r.segment.hosts} / {r.growth_pct}%",
                        f"{r.v4 or '—'}\n{r.v4.netmask if r.v4 else '—'}",
                        ", ".join(map(str, r.reservation.gateway_addresses))
                        if r.reservation
                        else "—",
                        f"Statyczne: {address_range(r.reservation.static_range) or '—'}\n"
                        f"DHCP: {address_range(r.reservation.dhcp_range) or '—'}"
                        if r.reservation
                        else "—",
                    )
                    for r in rows
                ),
                (4.3, 2.4, 5, 4.3, 10.3),
            )
            _table(
                document,
                ("Segment", "IPv6", "Brama i RA", "Końce i reverse DNS"),
                (
                    (
                        r.segment.name,
                        r.v6,
                        f"{r.info6.gateway or '—'}\n{r.info6.ra_mode or '—'}" if r.info6 else "—",
                        "\n".join((*map(str, r.info6.endpoints), r.info6.reverse_zone or ""))
                        if r.info6
                        else "—",
                    )
                    for r in rows
                ),
                (4.3, 6, 6, 10),
            )
            for r in rows:
                if r.segment.notes:
                    document.add_paragraph(office_text(r.segment.name + ": " + r.segment.notes))
    document.add_heading("Infrastruktura IPv6", 1)
    _table(
        document,
        ("Lokalizacja", "Blok", "Rezerwa", "P2P", "Loopback"),
        (
            (
                names.get(p.site_id, p.site_id),
                p.network,
                p.reserved_pool,
                p.p2p_pool,
                p.loopback_pool,
            )
            for p in report.plan.v6_infrastructure
        ),
        (4.3, 5.5, 5.5, 5.5, 5.5),
    )
    if project.active_blocked_subnets:
        document.add_heading("Zajęte podsieci — blokady", 1)
        _table(
            document,
            ("Pula", "CIDR", "Opis"),
            (
                (names.get(b.parent_id or "", project.name), b.network, b.note)
                for b in project.active_blocked_subnets
            ),
            (8, 8, 10.3),
        )
    document.add_heading("Wolne bloki", 1)
    _table(
        document,
        ("Rodzic", "Rodzaj puli", "CIDR"),
        (
            (names.get(b.parent_id or "", project.name), b.pool_kind or "IPv4", b.network)
            for b in report.plan.free_blocks
        ),
        (8, 6, 12.3),
    )
    document.add_heading("Walidacja", 1)
    _table(
        document,
        ("Poziom", "Element", "Komunikat"),
        (
            (f.level, names.get(f.node_id or "", project.name), format_finding(f, project))
            for f in report.plan.findings
        ),
        (2.5, 4.5, 19.3),
    )
    footer = section.footer.paragraphs[0]
    footer.add_run("Address Plan · " + report.status + " · strona ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "Address Plan"
    document.core_properties.title = office_text(project.name)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()

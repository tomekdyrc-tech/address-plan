"""Round-trip, safety and real HTTP/CLI contracts for Office reports."""

import subprocess
import sys
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from starlette.testclient import TestClient

from adresplan.application.exporting import export_bytes, save_document
from adresplan.application.project_json import ProjectJSONError, load_project
from adresplan.core.freeze import freeze
from adresplan.core.models import Project, Segment
from adresplan.core.plan import build_plan
from adresplan.export.csv_ import report_records
from adresplan.export.office import XLSX_COLUMNS, office_text
from adresplan.web.app import create_app
from adresplan.web.forms import edited_project

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("name", ["church-av", "budynek", "t5-ipv6", "t3-no-space"])
def test_xlsx_roundtrip_matches_source(name):
    project = load_project(ROOT / "examples" / f"{name}.json").project
    raw = export_bytes(project, "xlsx")
    book = load_workbook(BytesIO(raw), data_only=False)
    assert book.sheetnames == ["Plan", "Podsumowanie", "Wolne bloki", "Walidacja"]
    rows = list(book["Plan"].values)
    assert rows[0][0] == project.name
    assert (
        "NIEKOMPLETNY" in rows[1][0]
        if not build_plan(project).is_complete
        else "KOMPLETNY" in rows[1][0]
    )
    assert rows[3] == XLSX_COLUMNS
    for actual, record in zip(rows[4:], report_records(project), strict=True):
        assert actual == tuple(office_text(record.get(key, "")) for key in XLSX_COLUMNS)
    for sheet in book:
        assert sheet.freeze_panes
        assert sheet.auto_filter.ref
        assert sheet.print_title_rows == "$1:$4"
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    # The /48 count must never be rounded through an Excel floating-point number.
    for row in book["Podsumowanie"]:
        for cell in row:
            if cell.value:
                assert isinstance(cell.value, str)
    book.close()


def test_xlsx_rejects_oversized_cells_without_silent_truncation():
    project = load_project(ROOT / "examples/church-av.json").project
    segment = next(n for n in project.nodes if isinstance(n, Segment))
    oversized = replace(
        project,
        nodes=tuple(
            replace(n, notes="x" * 32768) if n.id == segment.id else n for n in project.nodes
        ),
    )
    with pytest.raises(ProjectJSONError, match="32767"):
        export_bytes(oversized, "xlsx")


@pytest.mark.parametrize("name", ["church-av", "budynek", "t5-ipv6", "t3-no-space"])
def test_docx_roundtrip_contains_plan_and_diagnostics(name):
    project = load_project(ROOT / "examples" / f"{name}.json").project
    raw = export_bytes(project, "docx")
    doc = Document(BytesIO(raw))
    text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )
    assert project.name in text
    assert (
        "Status planu: " + ("KOMPLETNY" if build_plan(project).is_complete else "NIEKOMPLETNY")
        in text
    )
    for allocation in build_plan(project).v4_allocations + build_plan(project).v6_allocations:
        assert str(allocation.network) in text
    for node in project.nodes:
        assert node.name in text
        if node.notes:
            assert node.notes in text
    assert all(t.rows[0]._tr.xpath("./w:trPr/w:tblHeader") for t in doc.tables)
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.sections[0].page_width > doc.sections[0].page_height
    with ZipFile(BytesIO(raw)) as archive:
        assert not any("vba" in name.lower() for name in archive.namelist())


@pytest.mark.parametrize(
    "text",
    [
        '=HYPERLINK("https://example.invalid","x")',
        "+SUM(A1)",
        "-1+1",
        "@SUM(A1)",
        "<script>bad</script>",
        "\t=1+1",
        "tekst\x00kontrolny",
        "ąćęłńóśźż\nDrugi wiersz",
    ],
)
def test_office_text_is_literal_and_xml_safe(text):
    project = load_project(ROOT / "examples/church-av.json").project
    segment = next(n for n in project.nodes if isinstance(n, Segment))
    project = replace(
        project,
        name=text,
        nodes=tuple(
            replace(n, name=text, notes=text) if n.id == segment.id else n for n in project.nodes
        ),
    )
    book = load_workbook(BytesIO(export_bytes(project, "xlsx")))
    assert book["Plan"]["A1"].value == office_text(text)
    assert book["Plan"]["A1"].data_type == "s"
    assert all(c.hyperlink is None and c.data_type != "f" for s in book for r in s for c in r)
    book.close()
    document = Document(BytesIO(export_bytes(project, "docx")))
    assert document.paragraphs[1].text == office_text(text)


@pytest.mark.parametrize("format", ["xlsx", "docx"])
def test_binary_atomic_write_invalid_and_empty_project(format, tmp_path, monkeypatch):
    project = load_project(ROOT / "examples/budynek.json").project
    empty = Project(id="empty", name="Pusty projekt", v4_pool=project.v4_pool)
    assert export_bytes(empty, format).startswith(b"PK")
    with pytest.raises(ProjectJSONError):
        export_bytes(replace(project, name=""), format)
    target = tmp_path / f"plan.{format}"
    raw = export_bytes(project, format)
    save_document(target, raw, format)
    assert target.read_bytes() == raw
    monkeypatch.setattr(
        "adresplan.application.files.os.replace", lambda *_: (_ for _ in ()).throw(OSError("test"))
    )
    with pytest.raises(ProjectJSONError):
        save_document(target, b"replacement", format)
    assert target.read_bytes() == raw


@pytest.mark.parametrize("format", ["xlsx", "docx"])
def test_cli_office_requires_file_and_marks_partial(format, tmp_path):
    target = tmp_path / f"report.{format}"
    args = [
        sys.executable,
        "-m",
        "adresplan",
        "export",
        str(ROOT / "examples/t3-no-space.json"),
        "--format",
        format,
    ]
    missing = subprocess.run(args, capture_output=True)
    assert missing.returncode == 2 and not missing.stdout
    result = subprocess.run([*args, "--output", str(target)], capture_output=True)
    assert result.returncode == 1
    assert target.read_bytes().startswith(b"PK")


def test_segment_copy_clears_both_pins_without_modifying_source():
    project = load_project(ROOT / "examples/church-av.json").project
    project = freeze(project, build_plan(project))
    segment = next(n for n in project.nodes if isinstance(n, Segment))
    copied = edited_project(project, {"action": "duplicate:" + segment.id})
    assert len(copied.nodes) == len(project.nodes) + 1
    new = next(n for n in copied.nodes if n.id not in {p.id for p in project.nodes})
    assert new.parent_id == segment.parent_id
    assert new.name == segment.name + " — kopia"
    assert new.pinned_v4 is None and new.pinned_v6 is None
    assert new.vlan_id == segment.vlan_id and new.hosts == segment.hosts
    assert next(n for n in copied.nodes if n.id == segment.id) == segment
    with pytest.raises(ProjectJSONError):
        edited_project(project, {"action": "duplicate:" + project.nodes[0].id})


@pytest.mark.parametrize("format", ["xlsx", "docx"])
def test_office_http_preview_download_and_print(format, tmp_path):
    app = create_app(tmp_path / "web.db")
    project = load_project(ROOT / "examples/t3-no-space.json").project
    saved = app.state.store.create(project)
    with TestClient(app, base_url="http://localhost") as client:
        preview = client.get(f"/p/{saved.key}/export?format={format}")
        assert preview.status_code == 200 and "PLAN NIEKOMPLETNY" in preview.text
        response = client.get(f"/api/projects/{saved.key}/export.{format}")
        assert response.status_code == 200
        assert response.headers["X-Plan-Complete"] == "false"
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert response.content.startswith(b"PK")
        printed = client.get(f"/p/{saved.key}/print")
        assert printed.status_code == 200 and "PLAN NIEKOMPLETNY" in printed.text
        assert 'id="project-form"' not in printed.text
        token = client.get("/api/session").json()["csrf_token"]
        result = client.post(
            f"/p/{saved.key}/save",
            data={
                "revision": 1,
                "action": "duplicate:"
                + next(n.id for n in project.nodes if isinstance(n, Segment)),
            },
            headers={"X-CSRF-Token": token, "HX-Request": "true"},
        )
        assert result.status_code == 200 and result.headers["HX-Retarget"] == "#workspace"
        assert len(app.state.store.get(saved.key).project.nodes) == len(project.nodes) + 1

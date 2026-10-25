"""User workflows for library management, filtering, backup and planning review."""

import sqlite3
from dataclasses import replace
from pathlib import Path

import pytest
from starlette.testclient import TestClient

from adresplan.application.project_json import load_project
from adresplan.application.templates import building_template
from adresplan.store.sqlite import MissingProject, ProjectStore
from adresplan.web.app import create_app


@pytest.fixture
def improvements(tmp_path):
    app = create_app(tmp_path / "library.db")
    with TestClient(app, base_url="http://localhost") as client:
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
        project = load_project(Path(__file__).parents[1] / "examples/church-av.json").project
        yield client, app.state.store, app.state.store.create(project)


def test_template_rename_replace_trash_restore_and_conflict(improvements):
    client, store, saved = improvements
    building = next(n for n in saved.project.nodes if n.kind == "building")
    key = store.save_template("Wzorzec", building_template(saved.project, building.id))
    assert client.get(f"/templates/{key}").status_code == 200
    response = client.post(
        f"/templates/{key}", data={"revision": 1, "name": "Oddział", "action": "rename"}
    )
    assert response.status_code == 200
    assert store.template_details(key)["name"] == "Oddział"
    response = client.post(f"/templates/{key}", data={"revision": 1, "name": "Stara karta"})
    assert response.status_code == 409
    response = client.post(
        f"/templates/{key}",
        data={
            "revision": 2,
            "action": "replace",
            "source_project": saved.key,
            "source_revision": saved.revision,
            "source_node": building.id,
        },
    )
    assert response.status_code == 200
    assert store.get(saved.key) == saved
    response = client.post(f"/templates/{key}", data={"revision": 3, "action": "archive"})
    assert response.status_code == 200 and not store.templates()
    with pytest.raises(MissingProject):
        store.template(key)
    assert store.templates(archived=True)[0]["key"] == key
    assert (
        client.post(f"/templates/{key}", data={"revision": 4, "action": "restore"}).status_code
        == 200
    )
    assert store.templates()[0]["key"] == key


def test_full_backup_retains_history_and_archived_templates(improvements, tmp_path):
    client, store, saved = improvements
    store.save(saved.key, replace(saved.project, name="Nowsza wersja"), saved.revision)
    building = next(n for n in saved.project.nodes if n.kind == "building")
    key = store.save_template("Archiwum", building_template(saved.project, building.id))
    store.update_template(key, 1, archived=True)
    response = client.get("/library/backup.sqlite3")
    assert response.status_code == 200 and response.content.startswith(b"SQLite format 3")
    restored_path = tmp_path / "restored.sqlite3"
    restored_path.write_bytes(response.content)
    restored = ProjectStore(restored_path)
    assert restored.get(saved.key) == store.get(saved.key)
    assert restored.history(saved.key) == store.history(saved.key)
    assert restored.template_details(key) == store.template_details(key)
    assert restored.templates(archived=True) == store.templates(archived=True)


def test_database_v2_migration_preserves_library(improvements):
    _, store, saved = improvements
    building = next(n for n in saved.project.nodes if n.kind == "building")
    key = store.save_template("Wzorzec", building_template(saved.project, building.id))
    with sqlite3.connect(store.path) as db:
        db.execute("ALTER TABLE building_templates DROP COLUMN archived")
        db.execute("ALTER TABLE building_templates DROP COLUMN revision")
        db.execute("PRAGMA user_version=2")
    migrated = ProjectStore(store.path)
    assert migrated.get(saved.key) == saved
    assert migrated.template_details(key)["revision"] == 1
    assert not migrated.template_details(key)["archived"]
    assert ProjectStore(store.path).template(key) == migrated.template(key)


def test_filtering_does_not_remove_inputs_or_mutate_project(improvements):
    client, store, saved = improvements
    response = client.get(f"/p/{saved.key}")
    assert response.status_code == 200
    assert 'id="view-filters"' in response.text
    assert response.text.index('id="view-filters"') < response.text.index('id="project-form"')
    assert response.text.count("data-filter-node=") == len(saved.project.nodes)
    for index, node in enumerate(saved.project.nodes):
        assert f'name="n{index}.name"' in response.text
        assert f'data-filter-result="{node.id}"' in response.text
    assert store.get(saved.key) == saved


def draft_from_html(text):
    from html.parser import HTMLParser

    class DraftParser(HTMLParser):
        value = None

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if tag == "input" and attrs.get("id") == "draft-document":
                self.value = attrs["value"]

    parser = DraftParser()
    parser.feed(text)
    assert parser.value
    return parser.value


def test_preview_structural_changes_stays_unsaved_then_commits_once(improvements):
    import json

    client, store, saved = improvements
    client.headers["HX-Request"] = "true"
    building = next(n for n in saved.project.nodes if n.kind == "building")
    response = client.post(
        f"/p/{saved.key}/save",
        data={
            "revision": saved.revision,
            "planning_mode": "review",
            "action": "add",
            "new_kind": "segment",
            "new_name": "Nowa sieć",
            "new_parent": building.id,
            "new_hosts": 12,
        },
    )
    assert response.status_code == 200
    assert response.headers["X-Plan-Preview"] == "1"
    assert store.get(saved.key) == saved
    draft = draft_from_html(response.text)
    data = json.loads(draft)
    assert len(data["project"]["nodes"]) == len(saved.project.nodes) + 1
    index = len(saved.project.nodes)
    response = client.post(
        f"/p/{saved.key}/save",
        data={
            "revision": saved.revision,
            "planning_mode": "review",
            "draft_document": draft,
            f"n{index}.hosts": "40",
        },
    )
    assert response.status_code == 200
    draft2 = draft_from_html(response.text)
    assert json.loads(draft2)["project"]["nodes"][-1]["hosts"] == 40
    assert store.get(saved.key) == saved
    client.headers.pop("HX-Request")
    response = client.post(
        f"/p/{saved.key}/apply-preview",
        data={"revision": saved.revision, "draft_document": draft2},
        follow_redirects=False,
    )
    assert response.status_code == 303
    applied = store.get(saved.key)
    assert applied.revision == saved.revision + 1
    assert applied.project.nodes[-1].hosts == 40
    assert len(store.history(saved.key)) == 2


def test_preview_stale_commit_and_foreign_project_are_rejected(improvements):
    from adresplan.application.project_json import dumps_project

    client, store, saved = improvements
    changed = replace(saved.project, name="Podgląd")
    draft = dumps_project(changed)
    store.save(saved.key, replace(saved.project, name="Nowszy zapis"), 1)
    assert (
        client.post(
            f"/p/{saved.key}/apply-preview", data={"revision": 1, "draft_document": draft}
        ).status_code
        == 409
    )
    other = dumps_project(replace(changed, id="other"))
    assert (
        client.post(
            f"/p/{saved.key}/save",
            data={"revision": 2, "planning_mode": "review", "draft_document": other},
        ).status_code
        == 422
    )
    assert store.get(saved.key).project.name == "Nowszy zapis"


def test_calculator_creates_project_without_double_counting(improvements):
    from adresplan.core.models import GatewayPosition
    from adresplan.core.plan import build_plan

    client, store, _ = improvements
    response = client.post(
        "/calc",
        data={
            "text": "192.168.8.0/24",
            "mode": "vlsm",
            "subnet_name": ["Biuro", "Goście"],
            "subnet_hosts": ["50", "20"],
            "action": "project",
            "project_name": "Z kalkulatora",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    project = store.get(response.headers["location"].rsplit("/", 1)[-1]).project
    assert project.settings.growth_pct == 0 and project.settings.gateway is GatewayPosition.NONE
    assert project.settings.gateway_addrs == 0 and not project.uses_ipv6
    assert {str(n.pinned_v4) for n in project.nodes if n.kind == "segment"} == {
        "192.168.8.0/26",
        "192.168.8.64/27",
    }
    assert build_plan(project).is_complete
    before = len(store.list())
    assert (
        client.post(
            "/calc",
            data={
                "text": "192.168.8.0/30",
                "mode": "vlsm",
                "subnet_name": "Za duża",
                "subnet_hosts": "50",
                "action": "project",
            },
        ).status_code
        == 422
    )
    assert len(store.list()) == before


def test_explanation_uses_same_sizing_rules(improvements):
    from adresplan.core.models import GatewayPosition, Segment, SegmentRole
    from adresplan.web.explanations import explain_segment

    _, _, saved = improvements
    project = replace(
        saved.project, settings=replace(saved.project.settings, growth_pct=25, gateway_addrs=1)
    )
    segment = Segment(id="x", parent_id="b", name="Test", hosts=20)
    text = explain_segment(project, segment)
    assert "20 urządzeń + 5 zapasu (25%) + 1" in text and "26 potrzebnych" in text and "/27" in text
    segment = replace(segment, gateway=GatewayPosition.NONE, gateway_addrs=0, static_count=4)
    assert "29 potrzebnych" in explain_segment(project, segment)
    p2p = replace(segment, role=SegmentRole.P2P, hosts=2, static_count=0)
    assert "/31" in explain_segment(project, p2p)


@pytest.mark.parametrize("format", ["md", "xlsx", "docx"])
def test_report_profiles_preserve_metadata_and_literal_text(improvements, format):
    from io import BytesIO

    from docx import Document
    from openpyxl import load_workbook

    from adresplan.application.exporting import export_bytes
    from adresplan.application.project_json import dumps_project, loads_project

    _, _, saved = improvements
    project = replace(
        saved.project,
        author="Autor",
        organization="=1+1",
        description="Opis projektu",
        document_date="2026-10-16",
    )
    assert loads_project(dumps_project(project)).project == project
    for profile in ("full", "brief"):
        content = export_bytes(project, format, profile)
        if format == "md":
            text = content.decode("utf-8")
            assert "Autor" in text and "Opis projektu" in text
        elif format == "xlsx":
            book = load_workbook(BytesIO(content))
            cells = [c for sheet in book for row in sheet for c in row]
            assert any(c.value == "=1+1" and c.data_type == "s" for c in cells)
        else:
            document = Document(BytesIO(content))
            assert any("Opis projektu" in p.text for p in document.paragraphs)


def test_profile_http_and_print(improvements):
    client, _, saved = improvements
    for format in ("md", "xlsx", "docx"):
        assert (
            client.get(f"/api/projects/{saved.key}/export.{format}?profile=brief").status_code
            == 200
        )
    response = client.get(f"/p/{saved.key}/print?profile=brief")
    assert response.status_code == 200 and "Raport skrócony" in response.text
    assert "Szczegóły techniczne" not in response.text


def test_large_project_form_saves_all_segments(improvements):
    from test_web import form

    from adresplan.core.models import Segment

    client, store, saved = improvements
    building = next(n for n in saved.project.nodes if n.kind == "building")
    segments = tuple(
        Segment(id=f"audit-{i}", parent_id=building.id, name=f"Segment {i}", hosts=5)
        for i in range(100)
    )
    project = replace(
        saved.project,
        nodes=tuple(n for n in saved.project.nodes if n.kind != "segment") + segments,
        v6_active=False,
    )
    saved = store.save(saved.key, project, saved.revision)
    values = form(client, saved.key)
    assert len(values) > 1000
    values["p.name"] = "Duży projekt po edycji"
    values[f"n{len(project.nodes) - 1}.hosts"] = "42"
    response = client.post(f"/p/{saved.key}/save", data=values)
    assert response.status_code == 200
    updated = store.get(saved.key)
    assert updated.project.name == values["p.name"]
    assert updated.project.nodes[-1].hosts == 42
    assert len(updated.project.nodes) == len(project.nodes)


def test_autosave_refreshes_filter_scope_cidr_and_sizing(improvements):
    from html.parser import HTMLParser

    from adresplan.core.plan import build_plan
    from adresplan.web.explanations import explain_segment

    client, store, saved = improvements
    index, segment = next((i, n) for i, n in enumerate(saved.project.nodes) if n.kind == "segment")
    response = client.post(
        f"/p/{saved.key}/save",
        data={
            "revision": saved.revision,
            f"n{index}.name": "Pracownia po zmianie",
            f"n{index}.hosts": "81",
            f"n{index}.vlan_id": "303",
        },
        headers={"HX-Request": "true"},
    )
    assert response.status_code == 200 and "HX-Retarget" not in response.headers

    class DerivedParser(HTMLParser):
        rows = None
        oob = False

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if attrs.get("id") == "editor-derived":
                self.oob = attrs.get("hx-swap-oob") == "outerHTML"
            if attrs.get("data-derived-node") == segment.id:
                self.rows = attrs

    parser = DerivedParser()
    parser.feed(response.text)
    changed = store.get(saved.key).project
    node = changed.nodes[index]
    cidr = next(a.network for a in build_plan(changed).v4_allocations if a.node_id == node.id)
    assert parser.oob and parser.rows
    assert "Pracownia po zmianie" in parser.rows["data-search"]
    assert "303" in parser.rows["data-search"] and str(cidr) in parser.rows["data-search"]
    assert node.parent_id in parser.rows["data-scope"]
    assert parser.rows["data-explanation"] == explain_segment(changed, node)


def test_invalid_template_name_does_not_partially_save_project(improvements):
    client, store, saved = improvements
    building = next(n for n in saved.project.nodes if n.kind == "building")
    response = client.post(
        f"/p/{saved.key}/save",
        data={
            "revision": saved.revision,
            "p.name": "Nie powinno się zapisać",
            "action": f"template-save:{building.id}",
            "template_name": "x" * 201,
        },
    )
    assert response.status_code == 422
    assert store.get(saved.key) == saved
    assert not store.templates()


def test_legacy_long_template_can_be_archived_and_restored(improvements):
    _, store, saved = improvements
    building = next(n for n in saved.project.nodes if n.kind == "building")
    key = store.save_template("Stara nazwa", building_template(saved.project, building.id))
    with sqlite3.connect(store.path) as db:
        db.execute("UPDATE building_templates SET name=? WHERE key=?", ("x" * 201, key))
    store.update_template(key, 1, archived=True)
    assert store.template_details(key)["archived"]
    store.update_template(key, 2, archived=False)
    assert not store.template_details(key)["archived"]
    store.update_template(key, 3, name="Nowa nazwa")
    assert store.template_details(key)["name"] == "Nowa nazwa"

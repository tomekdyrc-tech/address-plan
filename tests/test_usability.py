"""End-to-end data boundaries for the guided calculator and reusable locations."""

from dataclasses import replace
from ipaddress import IPv4Network

import pytest
from starlette.datastructures import FormData
from starlette.testclient import TestClient

from adresplan.application.templates import (
    apply_site_template,
    site_template,
    starter_templates,
)
from adresplan.web.app import create_app
from adresplan.web.calculator import calculator_context, friendly_error


def vlsm(pool="192.168.1.0/24", counts=(50, 25, 10)):
    return FormData(
        [
            ("text", pool),
            ("mode", "vlsm"),
            *[("subnet_name", f"Sieć {i}") for i in range(len(counts))],
            *[("subnet_hosts", str(c)) for c in counts],
        ]
    )


def test_vlsm_allocations_and_free_space():
    result = calculator_context(vlsm())
    networks = [r["subnet"].network for r in result["rows"]]
    assert [str(n) for n in networks] == ["192.168.1.0/26", "192.168.1.64/27", "192.168.1.96/28"]
    assert not result["unallocated"]
    all_blocks = [*networks, *result["free_blocks"]]
    assert sum(n.num_addresses for n in all_blocks) == 256
    assert all(not a.overlaps(b) for i, a in enumerate(all_blocks) for b in all_blocks[i + 1 :])


@pytest.mark.parametrize("count,prefix", [(1, 30), (2, 30), (62, 26), (63, 25), (254, 24)])
def test_vlsm_lan_host_boundaries(count, prefix):
    result = calculator_context(vlsm(counts=(count,)))
    assert result["rows"][0]["subnet"].network.prefixlen == prefix


@pytest.mark.parametrize("counts", [(), (0,), (-1,), (2**32,), ("abc",), tuple([2] * 129)])
def test_vlsm_rejects_invalid_requirements(counts):
    with pytest.raises(ValueError):
        calculator_context(vlsm(counts=counts))


def test_vlsm_ipv6_and_exhaustion_are_explicit():
    with pytest.raises(ValueError, match="IPv4"):
        calculator_context(vlsm("2001:db8::/48"))
    result = calculator_context(vlsm("192.168.1.0/26"))
    assert result["unallocated"] == ["Sieć 1", "Sieć 2"]


def test_split_mode_ignores_other_fields_and_errors_are_plain_polish():
    form = FormData({"text": "10.0.0.0/24", "mode": "count", "count": "4", "new_prefix": "28"})
    assert calculator_context(form)["preview"].new_prefix == 26
    with pytest.raises(ValueError) as error:
        calculator_context(FormData({"text": "10.0.0.0/24", "mode": "count", "count": "3"}))
    assert "Wpisz 4" in friendly_error(error.value)
    assert "--" not in friendly_error(error.value)
    preview = calculator_context(
        FormData({"text": "::/0", "mode": "prefix", "new_prefix": "128", "limit": "2"})
    )["preview"]
    assert preview.total_count == 2**128 and len(preview.subnets) == 2


@pytest.fixture
def ui(tmp_path):
    app = create_app(tmp_path / "usability.sqlite3")
    with TestClient(app, base_url="http://localhost") as client:
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
        yield client, app.state.store


def new_project(client):
    response = client.post(
        "/projects/new",
        data={
            "name": "Sieć testowa",
            "v4_pool": "10.0.0.0/16",
            "site_name": "Warszawa",
            "building_name": "Biuro",
            "start_template": "building|starter:office",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    return response.headers["location"].split("/")[-1]


def test_guided_creation_and_site_templates_survive_restart(ui):
    client, store = ui
    key = new_project(client)
    saved = store.get(key)
    assert len(saved.project.nodes) == 5
    site_id = saved.project.nodes[0].id
    response = client.post(
        f"/p/{key}/save", data={"revision": 1, "action": f"site-template-save:{site_id}"}
    )
    assert response.status_code == 200
    template_key = store.templates()[0]["key"]
    recipe = store.template(template_key)
    assert recipe["site"]["name"] == "Warszawa"
    response = client.post(
        f"/p/{key}/save",
        data={
            "revision": 1,
            "action": "site-template-apply",
            "site_template_key": template_key,
            "template_site_names": "Kraków\nGdańsk",
        },
    )
    assert response.status_code == 200
    project = store.get(key).project
    assert len(project.nodes) == 15
    assert len({n.id for n in project.nodes}) == 15
    assert all(n.pinned_v4 is None and n.pinned_v6 is None for n in project.nodes)
    assert [n.name for n in project.nodes if n.kind == "site"] == ["Warszawa", "Kraków", "Gdańsk"]
    assert store.get(key).revision == 2
    # New projects can also start with a previously saved whole location.
    response = client.post(
        "/projects/new",
        data={"name": "Inna sieć", "site_name": "Poznań", "start_template": f"site|{template_key}"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    copied = store.get(response.headers["location"].split("/")[-1]).project
    assert len(copied.nodes) == 5 and copied.nodes[0].name == "Poznań"
    assert not ({n.id for n in copied.nodes} & {n.id for n in project.nodes})
    assert template_key in client.get("/").text


def test_wrong_template_type_and_failed_batch_do_not_write(ui):
    client, store = ui
    key = new_project(client)
    project = store.get(key).project
    template_key = store.save_template("Oddział", site_template(project, project.nodes[0].id))
    for action, extra in [
        ("template-apply", {"template_key": template_key}),
        ("site-template-apply", {"site_template_key": template_key, "template_site_names": ""}),
    ]:
        response = client.post(f"/p/{key}/save", data={"revision": 1, "action": action, **extra})
        assert response.status_code == 422
        assert store.get(key).revision == 1


def test_calculator_html_and_manual_add(ui):
    client, store = ui
    response = client.post(
        "/calc",
        data={
            "text": "192.168.1.0/24",
            "mode": "vlsm",
            "subnet_name": "<script>x</script>",
            "subnet_hosts": "50",
        },
    )
    assert response.status_code == 200 and "192.168.1.0/26" in response.text
    assert "<script>x</script>" not in response.text
    assert "&lt;script&gt;" in response.text
    key = new_project(client)
    project = store.get(key).project
    building_id = next(n.id for n in project.nodes if n.kind == "building")
    response = client.post(
        f"/p/{key}/save",
        data={
            "revision": 1,
            "action": "add",
            "new_kind": "segment",
            "new_name": "Kamery",
            "new_parent": building_id,
            "new_hosts": "24",
            "new_vlan": "40",
        },
    )
    assert response.status_code == 200
    segment = store.get(key).project.nodes[-1]
    assert segment.hosts == 24 and segment.vlan_id == 40
    assert response.text.index('class="add-node"') < response.text.index('class="edit-table"')


def test_site_template_discards_pins_and_preserves_segment_settings(ui):
    client, store = ui
    project = store.get(new_project(client)).project
    nodes = tuple(replace(n, pinned_v4=IPv4Network("10.0.0.0/24")) for n in project.nodes)
    recipe = site_template(replace(project, nodes=nodes), nodes[0].id)
    copied = apply_site_template(project, recipe, "Oddział")
    assert all(n.pinned_v4 is None for n in copied.nodes[5:])
    assert copied.nodes[-1].hosts == project.nodes[-1].hosts
    assert copied.nodes[-1].growth_pct == project.settings.growth_pct


def test_new_projects_require_explicit_ipv6_choice(ui):
    from adresplan.core.plan import build_plan

    client, store = ui
    key = new_project(client)
    assert not store.get(key).project.uses_ipv6
    assert not build_plan(store.get(key).project).v6_allocations
    html = client.get(f"/p/{key}").text
    assert '<option value="no" selected>' in html
    assert "<th>IPv6 / RA</th>" not in html
    response = client.post(
        "/projects/new",
        data={
            "name": "Dual stack",
            "ipv6_usage": "yes",
            "start_template": "building|starter:office",
        },
        follow_redirects=False,
    )
    project = store.get(response.headers["location"].split("/")[-1]).project
    assert project.v6_prefix is not None and project.v6_kind == "ula"
    assert str(project.v6_prefix).startswith("fd")
    assert build_plan(project).v6_allocations


def test_ipv6_toggle_is_persistent_and_does_not_regenerate_on_save(ui):
    client, store = ui
    key = new_project(client)
    client.headers["HX-Request"] = "true"
    response = client.post(f"/p/{key}/save", data={"revision": 1, "ipv6_usage": "yes"})
    assert response.status_code == 200 and response.headers["HX-Retarget"] == "#workspace"
    first = store.get(key)
    assert first.project.v6_prefix is not None
    response = client.post(f"/p/{key}/save", data={"revision": first.revision, "ipv6_usage": "yes"})
    assert response.status_code == 200
    assert store.get(key) == first
    assert "<th>IPv6 / RA</th>" in client.get(f"/p/{key}").text
    response = client.post(
        f"/p/{key}/save",
        data={
            "revision": first.revision,
            "ipv6_usage": "no",
            "action": "ula",
            "p.v6_prefix": "bad",
            "s.v6_site_prefix": "bad",
        },
    )
    assert response.status_code == 200
    assert store.get(key).project.v6_prefix == first.project.v6_prefix
    assert not store.get(key).project.v6_active
    assert store.snapshot(key, first.revision)["project"]["v6_prefix"] == str(
        first.project.v6_prefix
    )


def test_disabling_ipv6_preserves_both_families(ui):
    from ipaddress import IPv6Network

    from adresplan.core.freeze import freeze
    from adresplan.core.models import BlockedSubnet
    from adresplan.core.plan import build_plan
    from adresplan.web.forms import edited_project

    client, store = ui
    project = store.get(new_project(client)).project
    project = replace(project, v6_prefix=IPv6Network("2001:db8:42::/48"), v6_active=True)
    project = freeze(project, build_plan(project))
    project = replace(
        project,
        blocked_subnets=(
            BlockedSubnet(network=IPv4Network("10.8.0.0/24")),
            BlockedSubnet(network=IPv6Network("2001:db8:42:200::/56")),
        ),
    )
    disabled = edited_project(project, {"ipv6_usage": "no"})
    assert disabled.v6_prefix == project.v6_prefix and not disabled.v6_active
    assert disabled.pinned_v6_infrastructure == project.pinned_v6_infrastructure
    assert [n.pinned_v6 for n in disabled.nodes] == [n.pinned_v6 for n in project.nodes]
    assert [n.pinned_v4 for n in disabled.nodes] == [n.pinned_v4 for n in project.nodes]
    assert disabled.blocked_subnets == project.blocked_subnets
    assert not build_plan(disabled).v6_allocations


def test_ipv6_cannot_be_added_by_block_or_template_when_disabled(ui):
    client, store = ui
    key = new_project(client)
    site = store.get(key).project.nodes[0].id
    response = client.post(
        f"/p/{key}/save",
        data={
            "revision": 1,
            "ipv6_usage": "no",
            "action": "block-add",
            "blocked_networks": "2001:db8::/64",
        },
    )
    assert response.status_code == 422 and store.get(key).revision == 1
    response = client.post(
        f"/p/{key}/save",
        data={
            "revision": 1,
            "ipv6_usage": "no",
            "action": "template-apply",
            "template_key": "starter:av",
            "template_site": site,
        },
    )
    assert response.status_code == 200
    assert store.get(key).project.v6_prefix is None
    response = client.post(f"/p/{key}/save", data={"revision": 2, "ipv6_usage": "maybe"})
    assert response.status_code == 422 and store.get(key).revision == 2


def test_starter_names_are_generic(ui):
    client, _ = ui
    html = client.get("/").text
    assert "Obiekt podstawowy" in html and "Obiekt rozbudowany" in html
    assert "realizacją AV" not in html and "Obiekt AV" not in html
    assert starter_templates()["av"]["segments"][1]["name"] == "Urządzenia i usługi"


@pytest.mark.parametrize("site_count", [2, 16])
def test_new_ipv6_hierarchy_fits_multiple_locations(ui, site_count):
    from adresplan.core.plan import build_plan

    client, store = ui
    key = new_project(client)
    source = store.get(key).project
    recipe = site_template(source, source.nodes[0].id)
    project = apply_site_template(
        source, recipe, "\n".join(f"Oddział {i}" for i in range(1, site_count))
    )
    from adresplan.web.forms import edited_project

    enabled = edited_project(project, {"ipv6_usage": "yes"})
    assert enabled.settings.v6_site_prefix == 52
    plan = build_plan(enabled)
    assert plan.is_complete
    assert len(plan.v6_allocations) == 5 * site_count
    assert (
        len(
            {
                a.network
                for a in plan.v6_allocations
                if any(n.id == a.node_id and n.kind == "site" for n in enabled.nodes)
            }
        )
        == site_count
    )


def test_existing_hierarchy_requires_explicit_fit_and_rejects_pins(ui):
    from ipaddress import IPv6Network

    from adresplan.application.ipv6_setup import fit_ipv6_hierarchy
    from adresplan.application.project_json import ProjectJSONError
    from adresplan.core.freeze import freeze
    from adresplan.core.plan import build_plan
    from adresplan.web.forms import edited_project

    client, store = ui
    source = store.get(new_project(client)).project
    project = replace(source, v6_prefix=IPv6Network("2001:db8:42::/48"), v6_active=True)
    project = apply_site_template(
        project, site_template(project, project.nodes[0].id), "Drugi oddział"
    )
    assert edited_project(project, {"ipv6_usage": "yes"}).settings.v6_site_prefix == 48
    assert not build_plan(project).is_complete
    fitted = fit_ipv6_hierarchy(project)
    assert build_plan(fitted).is_complete
    assert build_plan(fitted).v4_allocations == build_plan(project).v4_allocations
    frozen = freeze(fitted, build_plan(fitted))
    with pytest.raises(ProjectJSONError, match="przypięte"):
        fit_ipv6_hierarchy(frozen)
    with pytest.raises(ProjectJSONError, match="włącz"):
        fit_ipv6_hierarchy(replace(fitted, v6_active=False))


def test_fit_handles_more_than_sixteen_sites(ui):
    from adresplan.application.ipv6_setup import fit_ipv6_hierarchy
    from adresplan.core.plan import build_plan
    from adresplan.web.forms import edited_project

    client, store = ui
    source = store.get(new_project(client)).project
    project = apply_site_template(
        source,
        site_template(source, source.nodes[0].id),
        "\n".join(f"Oddział {i}" for i in range(16)),
    )
    project = edited_project(project, {"ipv6_usage": "yes"})
    assert not build_plan(project).is_complete
    fitted = fit_ipv6_hierarchy(project)
    assert fitted.settings.v6_site_prefix == 56
    assert fitted.settings.v6_building_prefix == 60
    assert build_plan(fitted).is_complete


def test_suspended_ipv6_survives_json_restart_freeze_and_reenable(ui):
    from ipaddress import IPv6Network

    from adresplan.application.project_json import dumps_project, loads_project
    from adresplan.core.freeze import freeze
    from adresplan.core.models import BlockedSubnet
    from adresplan.core.plan import build_plan
    from adresplan.store.sqlite import ProjectStore
    from adresplan.web.forms import edited_project

    client, store = ui
    source = store.get(new_project(client)).project
    source = edited_project(source, {"ipv6_usage": "yes"})
    source = replace(source, blocked_subnets=(BlockedSubnet(network=IPv6Network("fd00::/64")),))
    # An inactive block outside the pool is preserved, then validated on reactivation.
    suspended = edited_project(source, {"ipv6_usage": "no"})
    assert build_plan(suspended).is_complete
    assert loads_project(dumps_project(suspended)).project == suspended
    saved = store.create(suspended)
    assert ProjectStore(store.path).get(saved.key).project == suspended
    frozen4 = freeze(suspended, build_plan(suspended))
    assert frozen4.v6_prefix == source.v6_prefix
    assert frozen4.blocked_subnets == source.blocked_subnets
    assert not build_plan(edited_project(frozen4, {"ipv6_usage": "yes"})).is_complete


def test_suspended_frozen_ipv6_reactivates_at_same_addresses(ui):
    from adresplan.core.freeze import freeze
    from adresplan.core.plan import build_plan
    from adresplan.web.forms import edited_project

    client, store = ui
    enabled = edited_project(store.get(new_project(client)).project, {"ipv6_usage": "yes"})
    frozen = freeze(enabled, build_plan(enabled))
    off = edited_project(frozen, {"ipv6_usage": "no"})
    off = freeze(off, build_plan(off))
    on = edited_project(off, {"ipv6_usage": "yes"})
    assert on == frozen
    assert build_plan(on).v6_allocations == build_plan(frozen).v6_allocations
    assert on.pinned_v6_infrastructure == frozen.pinned_v6_infrastructure


def test_readable_findings_hide_internal_ids_and_link_to_fields(ui):
    from adresplan.core.models import Finding
    from adresplan.web.findings import present_finding

    client, store = ui
    key = new_project(client)
    project = store.get(key).project
    segment = project.nodes[-1]
    message = present_finding(Finding(code="insufficient_space", node_id=segment.id), project)
    assert "Warszawa → Biuro → Zarządzanie" in message.summary
    assert "dla 5 urządzeń" in message.summary
    assert segment.id not in message.summary and "insufficient_space" not in message.summary
    assert segment.id in message.technical
    assert message.field == "n4.hosts"
    html = client.get(f"/p/{key}").text
    assert "Przejdź do pola" in html
    assert "<summary>Informacje o planie" in html
    assert '<details class="validation"><summary>Informacje' in html


def test_suspend_resume_and_fit_http(ui):
    client, store = ui
    key = new_project(client)
    response = client.post(f"/p/{key}/save", data={"revision": 1, "ipv6_usage": "yes"})
    assert response.status_code == 200
    enabled = store.get(key)
    response = client.post(
        f"/p/{key}/save", data={"revision": enabled.revision, "ipv6_usage": "no"}
    )
    assert response.status_code == 200
    off = store.get(key)
    assert off.project.v6_prefix == enabled.project.v6_prefix and not off.project.v6_active
    assert "<th>IPv6 / RA</th>" not in response.text
    response = client.post(
        f"/p/{key}/save", data={"revision": off.revision, "ipv6_usage": "yes", "action": "ipv6-fit"}
    )
    assert response.status_code == 200
    assert store.get(key).project.v6_prefix == enabled.project.v6_prefix


@pytest.mark.parametrize("schema_version", [1, 2])
def test_legacy_json_preserves_ipv6_state(ui, schema_version):
    import json

    from adresplan.application.project_json import dumps_project, loads_project
    from adresplan.core.plan import build_plan
    from adresplan.web.forms import edited_project

    client, store = ui
    project = edited_project(store.get(new_project(client)).project, {"ipv6_usage": "yes"})
    document = json.loads(dumps_project(project))
    document["schema_version"] = schema_version
    del document["project"]["v6_active"]
    legacy = loads_project(json.dumps(document)).project
    assert legacy.v6_active and legacy.v6_prefix == project.v6_prefix
    assert build_plan(legacy).v6_allocations == build_plan(project).v6_allocations


def test_suspended_ipv6_exports_are_explicit(ui):
    from io import BytesIO

    from docx import Document
    from openpyxl import load_workbook

    from adresplan.application.exporting import export_bytes, export_document
    from adresplan.export.diagram import dumps_svg
    from adresplan.web.forms import edited_project

    client, store = ui
    project = edited_project(store.get(new_project(client)).project, {"ipv6_usage": "yes"})
    project = edited_project(project, {"ipv6_usage": "no"})
    assert "Wyłączone (konfiguracja zachowana)" in export_document(project, "md")
    assert "Wyłączone (konfiguracja zachowana)" in dumps_svg(project)
    workbook = load_workbook(BytesIO(export_bytes(project, "xlsx")))
    assert any(
        "konfiguracja zachowana" in str(cell.value)
        for sheet in workbook
        for row in sheet
        for cell in row
    )
    document = Document(BytesIO(export_bytes(project, "docx")))
    assert any("konfiguracja zachowana" in p.text for p in document.paragraphs)

"""Release 1.1: migration, historical evidence, templates and scoped exclusions."""

import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
from ipaddress import IPv4Network, IPv6Network
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from hypothesis import given
from hypothesis import strategies as st
from openpyxl import load_workbook
from starlette.testclient import TestClient

from adresplan.application.history import compare_snapshots
from adresplan.application.project_json import ProjectJSONError, dumps_project, loads_project
from adresplan.application.templates import apply_template, building_template
from adresplan.core.allocator import allocate
from adresplan.core.freeze import freeze
from adresplan.core.models import AllocationRequest, BlockedSubnet, Node, NodeKind, Project, Segment
from adresplan.core.plan import build_plan
from adresplan.export.diagram import dumps_svg
from adresplan.export.markdown import dumps_markdown
from adresplan.export.office import dumps_docx, dumps_xlsx
from adresplan.store.sqlite import MissingProject, ProjectStore, RevisionConflict, duplicate_project
from adresplan.web.app import create_app
from adresplan.web.forms import edited_project


@pytest.fixture
def simple():
    return Project(
        id="p",
        name="Oddział",
        v4_pool=IPv4Network("10.0.0.0/16"),
        nodes=(
            Node(id="s", name="Gdańsk", kind=NodeKind.SITE),
            Node(id="b", name="Biuro", kind=NodeKind.BUILDING, parent_id="s"),
            Segment(id="lan", name="Pracownicy", parent_id="b", hosts=20, vlan_id=10),
        ),
    )


def test_v1_database_migration_is_lossless_and_idempotent(tmp_path, simple):
    path = tmp_path / "old.sqlite3"
    document = json.loads(dumps_project(simple))
    document["schema_version"] = 1
    del document["project"]["blocked_subnets"]
    del document["plan"]
    with sqlite3.connect(path) as db:
        db.execute(
            "CREATE TABLE projects (key TEXT PRIMARY KEY, name TEXT NOT NULL, "
            "document TEXT NOT NULL, revision INTEGER NOT NULL, updated_at TEXT NOT NULL)"
        )
        db.execute(
            "INSERT INTO projects VALUES (?, ?, ?, ?, ?)",
            ("old", simple.name, json.dumps(document), 7, "2026-10-14"),
        )
        db.execute("PRAGMA user_version=1")
    for _ in range(2):
        store = ProjectStore(path)
        assert store.get("old").project == simple
        assert store.get("old").revision == 7
        assert [r.revision for r in store.history("old")] == [7]
        assert store.snapshot("old", 7)["plan"]["is_complete"]
        with sqlite3.connect(path) as db:
            assert json.loads(db.execute("SELECT document FROM projects").fetchone()[0]) == document
            assert db.execute("PRAGMA user_version").fetchone()[0] == 3
    with pytest.raises(MissingProject):
        store.snapshot("old", 1)


def test_failed_migration_rolls_back(tmp_path):
    path = tmp_path / "bad.sqlite3"
    with sqlite3.connect(path) as db:
        db.execute(
            "CREATE TABLE projects (key TEXT PRIMARY KEY, name TEXT, document TEXT, "
            "revision INTEGER, updated_at TEXT)"
        )
        db.execute("INSERT INTO projects VALUES ('bad', 'bad', '{}', 3, 'date')")
        db.execute("PRAGMA user_version=1")
    before = path.read_bytes()
    with pytest.raises(ProjectJSONError):
        ProjectStore(path)
    assert path.read_bytes() == before


def test_history_restore_conflicts_and_immutable_snapshots(tmp_path, simple):
    store = ProjectStore(tmp_path / "history.db")
    one = store.create(simple)
    snapshot = store.snapshot(one.key, 1)
    assert store.save(one.key, simple, 1) == one  # no phantom autosave revisions
    two = store.save(one.key, replace(simple, name="Nowa nazwa"), 1)
    with pytest.raises(RevisionConflict):
        store.restore(one.key, 1, 1)
    three = store.restore(one.key, 1, two.revision)
    assert three.project == simple and three.revision == 3
    assert store.snapshot(one.key, 1) == snapshot
    assert store.restore(one.key, 1, 3).revision == 4
    assert len(store.history(one.key)) == 4
    assert store.history(one.key, offset=2, limit=1)[0].revision == 2
    with pytest.raises(MissingProject):
        store.snapshot("other", 1)


def test_concurrent_edits_create_only_one_snapshot(tmp_path, simple):
    store = ProjectStore(tmp_path / "race.db")
    saved = store.create(simple)

    def change(name):
        try:
            store.save(saved.key, replace(simple, name=name), 1)
            return True
        except RevisionConflict:
            return False

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(change, ["A", "B"])) == [False, True]
    assert [r.revision for r in store.history(saved.key)] == [2, 1]


def test_diff_tracks_input_and_addresses_by_id(simple):
    old = json.loads(dumps_project(simple))
    changed = replace(
        simple,
        nodes=tuple(replace(n, hosts=90) if isinstance(n, Segment) else n for n in simple.nodes),
    )
    new = json.loads(dumps_project(changed))
    changes = compare_snapshots(old, new)
    assert any(c.field == "nodes.hosts" and c.before == "20" and c.after == "90" for c in changes)
    assert any(c.field == "v4_allocations.network" and "Pracownicy" in c.element for c in changes)
    assert compare_snapshots(old, old) == []
    # Historic results must be compared as recorded, without executing today's engine.
    new["plan"]["v4_allocations"][0]["network"] = "10.99.0.0/24"
    assert any(c.after == "10.99.0.0/24" for c in compare_snapshots(old, new))


@pytest.mark.parametrize("scope", [None, "s", "b"])
def test_v4_blocks_exclude_addresses_at_each_level(simple, scope):
    network = IPv4Network("10.0.0.0/24")
    project = replace(simple, blocked_subnets=(BlockedSubnet(network=network, parent_id=scope),))
    plan = build_plan(project)
    assert plan.is_complete
    nodes = {n.id: n for n in project.nodes}
    for allocation in plan.v4_allocations:
        if nodes[allocation.node_id].parent_id == scope:
            assert not allocation.network.overlaps(network)
    for block in plan.free_blocks:
        if block.parent_id == scope and block.network.version == 4:
            assert not block.network.overlaps(network)
    assert all(
        m.allocated_addresses + m.free_addresses == m.network.num_addresses for m in plan.v4_metrics
    )
    assert loads_project(dumps_project(project)).project == project
    assert build_plan(freeze(project, plan)) == plan


def test_overlapping_blocks_count_once(simple):
    blocks = tuple(
        BlockedSubnet(network=IPv4Network(n)) for n in ("10.0.0.0/24", "10.0.0.0/25", "10.0.0.0/24")
    )
    plan = build_plan(replace(simple, blocked_subnets=blocks))
    root = next(m for m in plan.v4_metrics if m.parent_id is None)
    site = next(a.network for a in plan.v4_allocations if a.node_id == "s")
    assert root.allocated_addresses == 256 + site.num_addresses


@pytest.mark.parametrize(
    "scope,network",
    [
        (None, "10.1.0.0/24"),
        ("missing", "10.0.0.0/24"),
        ("lan", "10.0.0.0/24"),
        (None, "2001:db8::/48"),
    ],
)
def test_invalid_blocks_are_incomplete(simple, scope, network):
    cls = IPv6Network if ":" in network else IPv4Network
    plan = build_plan(
        replace(simple, blocked_subnets=(BlockedSubnet(network=cls(network), parent_id=scope),))
    )
    assert not plan.is_complete
    assert any(f.code == "invalid_blocked_subnet" for f in plan.findings)


def test_frozen_conflict_never_moves_network(simple):
    frozen = freeze(simple, build_plan(simple))
    project = replace(
        frozen, blocked_subnets=(BlockedSubnet(network=IPv4Network("10.0.0.0/28"), parent_id="b"),)
    )
    plan = build_plan(project)
    assert any(f.code == "pinned_blocked" and f.node_id == "lan" for f in plan.findings)
    assert not any(a.node_id == "lan" for a in plan.v4_allocations)
    with pytest.raises(ValueError):
        freeze(project, plan)


@pytest.mark.parametrize(
    "scope,cidr",
    [(None, "2001:db8:abcd::/48"), ("site", "2001:db8:abcd::/56"), ("a", "2001:db8:abcd:100::/64")],
)
def test_v6_blocks_and_infrastructure_move(t5_project, scope, cidr):
    network = IPv6Network(cidr)
    project = replace(
        t5_project,
        v6_prefix=IPv6Network("2001:db8::/32"),
        blocked_subnets=(BlockedSubnet(network=network, parent_id=scope),),
    )
    plan = build_plan(project)
    assert plan.is_complete, plan.findings
    nodes = {n.id: n for n in project.nodes}
    for a in plan.v6_allocations:
        node = nodes[a.node_id]
        if node.parent_id == scope and not (
            isinstance(node, Segment) and node.role in ("p2p", "loopback")
        ):
            assert not a.network.overlaps(network)
    if scope == "site":
        assert not plan.v6_infrastructure[0].network.overlaps(network)
    assert build_plan(freeze(project, plan)) == plan


@given(st.sets(st.integers(min_value=0, max_value=255), max_size=40))
def test_allocator_exclusion_conservation(octets):
    pool = IPv4Network("10.0.0.0/24")
    exclusions = tuple(IPv4Network(f"10.0.0.{n}/32") for n in sorted(octets))
    result = allocate(
        pool, (AllocationRequest(node_id="one", name="one", prefix=28),), excluded=exclusions
    )
    assert (
        sum(a.network.num_addresses for a in result.allocations)
        + sum(b.num_addresses for b in result.free_blocks)
        + len(octets)
        == 256
    )
    assert all(not a.network.overlaps(b) for a in result.allocations for b in exclusions)


def test_building_clone_and_reusable_template(tmp_path, simple):
    original = freeze(simple, build_plan(simple))
    copied = edited_project(original, {"action": "duplicate:b"})
    assert copied.nodes[:3] == original.nodes
    added = copied.nodes[3:]
    assert len(added) == 2 and added[1].parent_id == added[0].id
    assert all(n.pinned_v4 is None and n.pinned_v6 is None for n in added)
    assert not {n.id for n in added} & {n.id for n in original.nodes}
    store = ProjectStore(tmp_path / "templates.db")
    recipe = building_template(original, "b")
    key = store.save_template("Biuro", recipe)
    restarted = ProjectStore(store.path)
    other = replace(
        simple, nodes=(simple.nodes[0],), settings=replace(simple.settings, growth_pct=99)
    )
    result = apply_template(other, restarted.template(key), "s", "Nowe biuro")
    assert result.nodes[-1].growth_pct == simple.settings.growth_pct
    assert build_plan(result).is_complete
    assert restarted.templates() == [{"key": key, "name": "Biuro"}]


def test_blocks_remap_on_project_copy_and_drop_on_subtree_delete(simple):
    project = replace(
        simple, blocked_subnets=(BlockedSubnet(network=IPv4Network("10.0.1.0/24"), parent_id="b"),)
    )
    copied = duplicate_project(project)
    assert copied.blocked_subnets[0].parent_id == copied.nodes[1].id
    assert edited_project(project, {"action": "delete:b"}).blocked_subnets == ()
    assert len(edited_project(project, {"action": "duplicate:b"}).blocked_subnets) == 1


def test_bulk_blocks_atomic_input_and_delete(simple):
    form = {
        "action": "block-add",
        "blocked_networks": "# comment\n10.0.8.0/24; Kamery\n10.0.9.0/24; Wi-Fi",
        "blocked_parent": "b",
    }
    project = edited_project(simple, form)
    assert len(project.blocked_subnets) == 2
    assert len(edited_project(project, form).blocked_subnets) == 2
    assert len(edited_project(project, {"action": "block-delete:0"}).blocked_subnets) == 1
    with pytest.raises(ProjectJSONError):
        edited_project(
            simple, {**form, "blocked_networks": form["blocked_networks"] + "\n10.0.1.1/24"}
        )


def test_svg_and_reports_include_safe_blocks(simple):
    project = replace(
        simple,
        name='<script>alert("x")</script>',
        blocked_subnets=(
            BlockedSubnet(network=IPv4Network("10.0.8.0/24"), note="<Kamery> & =1+1"),
        ),
    )
    svg = dumps_svg(project)
    root = ET.fromstring(svg)
    assert root.tag == "{http://www.w3.org/2000/svg}svg"
    assert not root.findall(".//{http://www.w3.org/2000/svg}script")
    assert "10.0.8.0/24" in svg and "&lt;Kamery&gt;" in svg
    assert "10.0.8.0/24" in dumps_markdown(project)
    book = load_workbook(BytesIO(dumps_xlsx(project)))
    assert book["Blokady"]["C5"].value == "<Kamery> & =1+1"
    assert book["Blokady"]["C5"].data_type == "s"
    doc = Document(BytesIO(dumps_docx(project)))
    assert any(
        "10.0.8.0/24" in cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )


def test_web_v11_workflow_and_csrf(tmp_path, simple):
    app = create_app(tmp_path / "web.db")
    saved = app.state.store.create(simple)
    key = saved.key
    with TestClient(app, base_url="http://localhost") as client:
        assert client.post(f"/p/{key}/history/1/restore", data={"revision": 1}).status_code == 403
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]

        def submit(action, **fields):
            return client.post(
                f"/p/{key}/save",
                data={"revision": app.state.store.get(key).revision, "action": action, **fields},
                follow_redirects=False,
            )

        assert submit("template-save:b").status_code == 303
        template = app.state.store.templates()[0]["key"]
        assert submit("template-apply", template_key=template, template_site="s").status_code == 303
        assert len(app.state.store.get(key).project.nodes) == 5
        assert submit("block-add", blocked_networks="10.0.8.0/24; Kamery").status_code == 303
        for route in ("", "/history", "/diagram", "/print"):
            assert client.get(f"/p/{key}{route}").status_code == 200
        entries = client.get(f"/api/projects/{key}/history").json()
        assert len(entries) == 3
        assert (
            client.get(f"/api/projects/{key}/history/1").json()["project"]["nodes"][2]["id"]
            == "lan"
        )
        svg = client.get(f"/api/projects/{key}/diagram.svg?download=true")
        assert "attachment" in svg.headers["content-disposition"]
        ET.fromstring(svg.text)
        assert client.post(f"/p/{key}/history/1/restore", data={"revision": 1}).status_code == 409
        response = client.post(
            f"/p/{key}/history/1/restore", data={"revision": 3}, follow_redirects=False
        )
        assert response.status_code == 303
        assert app.state.store.get(key).project == simple
        assert len(app.state.store.history(key)) == 4
        assert client.get(f"/p/{key}/history?before=999").status_code == 404


def test_old_examples_still_import():
    for path in (Path(__file__).resolve().parents[1] / "examples").glob("*.json"):
        old = loads_project(path.read_bytes())
        assert loads_project(dumps_project(old.project)) == old


def test_template_preserves_infrastructure_roles_and_no_gateway(t5_project, simple):
    from adresplan.core.models import GatewayPosition
    from adresplan.core.sizing import size_segment_v4

    result = apply_template(t5_project, building_template(t5_project, "a"), "site", "Kopia A")
    assert build_plan(result).is_complete
    for node in result.nodes[len(t5_project.nodes) :]:
        if isinstance(node, Segment):
            assert size_segment_v4(node, result.settings)
    no_gateway = replace(
        simple,
        nodes=tuple(
            replace(n, gateway=GatewayPosition.NONE) if isinstance(n, Segment) else n
            for n in simple.nodes
        ),
    )
    copied = apply_template(no_gateway, building_template(no_gateway, "b"), "s", "Bez bramy")
    assert copied.nodes[-1].gateway_addrs == 0
    assert build_plan(copied).is_complete


def test_ipv6_frozen_infrastructure_conflict(t5_project):
    frozen = freeze(t5_project, build_plan(t5_project))
    project = replace(
        frozen,
        blocked_subnets=(
            BlockedSubnet(network=frozen.pinned_v6_infrastructure[0].network, parent_id="site"),
        ),
    )
    plan = build_plan(project)
    assert not plan.is_complete
    assert any(f.code == "v6_pinned_blocked" for f in plan.findings)
    assert not plan.v6_infrastructure


def test_blocks_only_hierarchy_and_total_exhaustion(simple):
    project = replace(
        simple,
        nodes=simple.nodes[:2],
        blocked_subnets=(BlockedSubnet(network=IPv4Network("10.0.3.0/24"), parent_id="b"),),
    )
    plan = build_plan(project)
    assert plan.is_complete
    assert next(a.network for a in plan.v4_allocations if a.node_id == "b") == IPv4Network(
        "10.0.3.0/24"
    )
    exhausted = replace(simple, blocked_subnets=(BlockedSubnet(network=simple.v4_pool),))
    assert not build_plan(exhausted).is_complete
    assert not build_plan(exhausted).v4_allocations


def test_ipv6_reservations_are_not_multiplied_by_growth(simple):
    project = replace(
        simple,
        nodes=simple.nodes[:2],
        v6_prefix=IPv6Network("2001:db8::/48"),
        blocked_subnets=(
            BlockedSubnet(network=IPv6Network("2001:db8:0:1000::/52"), parent_id="b"),
        ),
    )
    assert build_plan(project).is_complete
    assert (
        next(a.network for a in build_plan(project).v6_allocations if a.node_id == "b").prefixlen
        == 52
    )
